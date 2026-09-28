#!/usr/bin/env python3
"""Fetch HiStock's Top 50 investment-trust (投信) net buy / net sell lists.

Headless port of `Financial Crawler.ipynb`: same page, same selectors, but
writes the workbook with openpyxl instead of xlwings, so it runs without a
desktop copy of Excel (e.g. inside Claude Code or a Managed Agent).

Usage:
    python3 fetch_trust_flows.py                       # fetch live, print JSON
    python3 fetch_trust_flows.py --xlsx out.xlsx       # also write Excel
    python3 fetch_trust_flows.py --html saved.html     # parse a saved page
    python3 fetch_trust_flows.py --history-dir history # also save history/YYYY-MM-DD.json
"""
import argparse
import json
import re
import os
import sys
from datetime import date, datetime, timedelta, timezone

from bs4 import BeautifulSoup as BS

URL = "https://histock.tw/stock/three.aspx?s=b"

# Column order on the page (and in the notebook's `fields` list).
FIELDS = ["code", "name", "price", "change_pct", "volume", "net"]
HEADERS_ZH = ["代號", "股票", "價格", "漲跌幅", "成交量", "買超/賣超"]

SECTIONS = {
    "buy": "tb-outline outline1",   # 投信買超
    "sell": "tb-outline outline2",  # 投信賣超
}

TAIPEI = timezone(timedelta(hours=8))

# The buy list's heading carries the trading date, e.g. "09-24 Top 50 投信買超排行".
DATE_RE = re.compile(r"(\d{2})-(\d{2})\s*Top\s*50\s*投信買超")

# TWSE/TPEx codes: 4-6 digits with an optional letter suffix (2330, 00878, 00632R).
CODE_RE = re.compile(r"^\d{4,6}[A-Z]?$")


def fetch_html(url=URL, timeout=20):
    import requests

    res = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=timeout)
    res.raise_for_status()
    return res.text


def to_number(text):
    """'1,234' -> 1234, '+2.35%' -> 2.35, '-' -> None. Non-numeric text is returned as-is."""
    cleaned = text.replace(",", "").replace("%", "").replace("+", "").strip()
    if cleaned in ("", "-", "--"):
        return None
    try:
        num = float(cleaned)
    except ValueError:
        return text
    return int(num) if num.is_integer() and "." not in cleaned else num


def parse_rows(div):
    rows = []
    for li in div.find_all("li"):
        cells = [span.text.strip() for span in li.find_all("span", {"class": "w58"})]
        # Data rows start with a stock code; this drops header rows and index
        # rows such as "TWOI 櫃檯指數", whose net figure is on a different scale.
        if len(cells) < len(FIELDS) or not CODE_RE.match(cells[0]):
            continue
        row = dict(zip(FIELDS, cells))
        for key in ("price", "change_pct", "volume", "net"):
            row[key] = to_number(row[key])
        rows.append(row)
    return rows


def parse(html):
    soup = BS(html, "html.parser")
    result = {}
    for side, css_class in SECTIONS.items():
        div = soup.find("div", {"class": css_class})
        if div is None:
            raise ValueError(
                f"Could not find <div class='{css_class}'> for the {side} list; "
                "the HiStock page layout may have changed."
            )
        result[side] = parse_rows(div)
    return result


def parse_data_date(html, today=None):
    """Trading date from the page heading as YYYY-MM-DD, or None if not found.

    The heading has no year: assume the current Taipei year, or last year if
    that would put the date in the future (a January run showing December data).
    """
    match = DATE_RE.search(BS(html, "html.parser").get_text(" ", strip=True))
    if not match:
        return None
    today = today or datetime.now(TAIPEI).date()
    month, day = int(match.group(1)), int(match.group(2))
    for year in (today.year, today.year - 1):
        try:
            candidate = date(year, month, day)
        except ValueError:  # e.g. 02-29 in a non-leap year
            continue
        if candidate <= today:
            return candidate.isoformat()
    return None


def write_xlsx(data, path):
    """Same layout as the notebook: buy list in A:F, sell list in G:L."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "sheet1"
    ws.cell(1, 1, "Top50投信買超").font = Font(bold=True)
    ws.cell(1, 7, "Top50投信賣超").font = Font(bold=True)
    for offset, side in ((1, "buy"), (7, "sell")):
        for i, header in enumerate(HEADERS_ZH):
            ws.cell(2, offset + i, header).font = Font(bold=True)
        for r, row in enumerate(data[side], start=3):
            for i, field in enumerate(FIELDS):
                ws.cell(r, offset + i, row[field])
    wb.save(path)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--html", help="parse a saved HTML file instead of fetching")
    ap.add_argument("--url", default=URL)
    ap.add_argument("--xlsx", help="also write an Excel workbook to this path")
    ap.add_argument("--history-dir", help="also save a snapshot to <dir>/<date>.json for streaks.py")
    ap.add_argument("--date", help="trading date for the snapshot (YYYY-MM-DD); default: the date on the page")
    args = ap.parse_args(argv)

    if args.html:
        with open(args.html, encoding="utf-8") as f:
            html = f.read()
        source = args.html
    else:
        html = fetch_html(args.url)
        source = args.url

    data = parse(html)
    out = {
        "source": source,
        "data_date": parse_data_date(html),
        "fetched_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "counts": {side: len(rows) for side, rows in data.items()},
        **data,
    }
    if args.xlsx:
        write_xlsx(data, args.xlsx)
        out["xlsx"] = args.xlsx
    if args.history_dir:
        date = args.date or out["data_date"]
        if date is None:
            sys.exit("Could not read the trading date from the page; pass --date YYYY-MM-DD.")
        os.makedirs(args.history_dir, exist_ok=True)
        path = os.path.join(args.history_dir, f"{date}.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump({"date": date, **out}, f, ensure_ascii=False, indent=2)
        out["snapshot"] = path

    json.dump(out, sys.stdout, ensure_ascii=False, indent=2)
    print()


if __name__ == "__main__":
    main()
