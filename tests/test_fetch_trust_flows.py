import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(__file__)
sys.path.insert(0, os.path.join(HERE, "..", "plugins", "tw-institutional-flows", "skills", "trust-flows", "scripts"))

import fetch_trust_flows as ftf  # noqa: E402

FIXTURE = os.path.join(HERE, "fixtures", "histock_sample.html")


class ParseTest(unittest.TestCase):
    def setUp(self):
        with open(FIXTURE, encoding="utf-8") as f:
            self.data = ftf.parse(f.read())

    def test_skips_header_rows(self):
        self.assertEqual(len(self.data["buy"]), 2)
        self.assertEqual(len(self.data["sell"]), 1)

    def test_coerces_numbers(self):
        tsmc = self.data["buy"][0]
        self.assertEqual(tsmc["code"], "2330")
        self.assertEqual(tsmc["price"], 1005.0)
        self.assertEqual(tsmc["change_pct"], 1.52)
        self.assertEqual(tsmc["volume"], 32118)
        self.assertEqual(tsmc["net"], 4210)
        self.assertEqual(self.data["sell"][0]["net"], -2870)

    def test_etf_code_kept_as_string(self):
        self.assertEqual(self.data["buy"][1]["code"], "00878")

    def test_missing_section_raises(self):
        with self.assertRaises(ValueError):
            ftf.parse("<html></html>")

    def test_write_xlsx_matches_notebook_layout(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "out.xlsx")
            ftf.write_xlsx(self.data, path)
            ws = load_workbook(path).active
            self.assertEqual(ws.cell(1, 1).value, "Top50投信買超")
            self.assertEqual(ws.cell(1, 7).value, "Top50投信賣超")
            self.assertEqual(ws.cell(2, 6).value, "買超/賣超")
            self.assertEqual(ws.cell(3, 1).value, "2330")
            self.assertEqual(ws.cell(3, 7).value, "2603")


if __name__ == "__main__":
    unittest.main()
